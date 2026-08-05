"""按工单导出累计产量源明细脚本测试。"""

from datetime import datetime, time
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.export_wrkorder_cumulative_details import (
    SOURCE_FIELDS,
    build_detail_query,
    export_wrkorder_details,
)


class FakeCursor:
    """模拟服务端游标，并记录查询参数。"""

    def __init__(self, rows: list[tuple]):
        """
        初始化模拟游标。

        Args:
            rows (list[tuple]): 查询返回的源数据行。
        """
        self.rows = rows
        self.description = [(field,) for field in SOURCE_FIELDS]
        self.sql = ""
        self.params: tuple = ()

    def __enter__(self):
        """
        进入游标上下文。

        Returns:
            FakeCursor: 当前模拟游标。
        """
        return self

    def __exit__(self, exc_type, exc_value, traceback):
        """
        退出游标上下文。

        Args:
            exc_type (type | None): 异常类型。
            exc_value (BaseException | None): 异常实例。
            traceback (TracebackType | None): 异常追踪信息。

        Returns:
            bool: 始终不抑制异常。
        """
        return False

    def execute(self, sql: str, params: tuple):
        """
        记录待执行 SQL 和参数。

        Args:
            sql (str): 参数化查询语句。
            params (tuple): 查询参数。
        """
        self.sql = sql
        self.params = params

    def __iter__(self):
        """
        迭代模拟查询结果。

        Returns:
            Iterator[tuple]: 源数据行迭代器。
        """
        return iter(self.rows)


class FakeConnection:
    """模拟只读数据库连接。"""

    def __init__(self, rows: list[tuple]):
        """
        初始化模拟连接。

        Args:
            rows (list[tuple]): 查询返回的源数据行。
        """
        self.fake_cursor = FakeCursor(rows)
        self.closed = False

    def cursor(self):
        """
        返回模拟游标。

        Returns:
            FakeCursor: 当前模拟游标。
        """
        return self.fake_cursor

    def close(self):
        """记录连接已关闭。"""
        self.closed = True


def test_build_detail_query_only_filters_exact_wrkorder():
    """查询只能按完整工单号过滤，不得汇总或添加业务过滤。"""
    sql = build_detail_query()
    normalized_sql = " ".join(sql.split()).lower()

    assert "where `wrkorder` = %s" in normalized_sql
    assert "group by" not in normalized_sql
    assert "sum(" not in normalized_sql
    assert "regdate >=" not in normalized_sql
    assert "flow in" not in normalized_sql
    assert not {"insert", "update", "delete", "replace"} & set(normalized_sql.split())


def test_export_writes_all_source_rows_without_processing(tmp_path: Path):
    """导出应逐行保留生产表返回的全部源字段和源值。"""
    source_row = (
        "T001",
        1,
        "BU1211",
        8,
        70,
        25,
        1001,
        datetime(2026, 8, 5),
        datetime(1899, 12, 30, 14, 30),
        "R001",
        "SO3-L3B",
        "PO001",
        12,
        "EST",
        8,
        "M",
        "BLACK",
        "L",
        "S001",
        "ST01",
    )
    connection = FakeConnection([source_row])

    output_path, row_count = export_wrkorder_details(
        wrkorder="BU1211",
        output_dir=tmp_path,
        db_config={"host": "example.invalid"},
        connection_factory=lambda **_kwargs: connection,
    )

    assert row_count == 1
    assert output_path.exists()
    assert connection.closed is True
    assert connection.fake_cursor.params == ("BU1211",)

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()

    assert rows[0] == SOURCE_FIELDS
    expected_row = source_row[:8] + (time(14, 30),) + source_row[9:]
    assert rows[1] == expected_row


def test_export_rejects_empty_wrkorder_before_connecting(tmp_path: Path):
    """空工单号必须在连接生产库前终止。"""
    connected = False

    def connection_factory(**_kwargs):
        """
        记录不应发生的连接动作。

        Args:
            **_kwargs: 模拟连接参数。
        """
        nonlocal connected
        connected = True
        raise AssertionError("不应连接数据库")

    with pytest.raises(ValueError, match="WRKORDER 不能为空"):
        export_wrkorder_details(
            wrkorder="  ",
            output_dir=tmp_path,
            db_config={},
            connection_factory=connection_factory,
        )

    assert connected is False
