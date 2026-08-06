"""按工单号从生产库导出累计产量源数据明细。"""

import re
import sys
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

import pymysql
from loguru import logger
from openpyxl import Workbook
from pymysql.cursors import SSCursor

# ======
# 查询模式配置
EXPORT_MODE = "detail"  # detail：单工单完整明细；compact：多工单精简汇总
WRKORDER = "BU1165"
WRKORDERS = (
    "BU0724",
    "BU0730",
    "BU0750",
    "BU0751",
    "BU0796",
    "BU0994",
    "BU1022",
    "BU1072",
    "BU1073",
    "BU1081",
    "BU1082",
    "BU1083",
    "BU1087",
    "BU1089",
    "BU1090",
    "BU1129",
    "BU1130",
    "BU1143",
    "BU1148",
    "BU1160",
    "BU1161",
    "BU1162",
    "BU1163",
    "BU1165",
    "BU1165A",
    "BU1166",
    "BU1166A",
    "BU1179",
    "BU1180",
    "BU1185",
    "BU1190",
    "BU1191",
    "BU1194",
    "BU1195",
    "BU1197",
    "BU1198",
    "BU1202",
    "BU1203",
    "BU1203A",
    "BU1205",
    "BU1210",
    "BU1211",
    "BU1217",
    "BU1218",
    "BU1220",
    "BU1229",
    "BU1237",
    "BU1242",
    "BU1243",
    "BU1245",
    "BU1246",
    "BU1255",
    "BU1256",
    "BU1257",
    "BU1258",
    "BU1259",
    "BU1260",
    "BU1261",
    "BU1262",
    "BU1268",
    "BU1269",
    "BU1270",
    "BU1271",
    "BU1272",
    "BU1273",
    "BU1276",
    "BU1277",
    "BU1284",
    "BU1285",
    "BU1289",
    "BU1290",
    "BU1291",
    "BU1292",
    "BU1293",
    "BU1294",
    "BU1295",
    "BU1300",
    "BU1310",
    "BU1311",
    "BU1312",
    "BU1313",
    "BU1314",
    "BU1317",
    "BU1318",
    "BU1319",
    "BU1320",
    "BU1328",
    "BU1339",
    "BU1342",
    "BU1377",
    "BU1388",
    "BU1389",
)
COMPACT_STEP_NOS = (1, 3, 6)

# ======
# 文件路径配置
OUTPUT_DIR = Path(__file__).resolve().parent / "output"
LOG_DIR = Path(__file__).resolve().parent.parent / "logs"
ENV_FILE_CANDIDATES = (
    Path(__file__).resolve().parent / "export.env",
    Path.cwd() / "export.env",
    Path(__file__).resolve().parent.parent / "iwork" / ".env",
)

# ======
# 生产数据库连接配置
IWORK_DB_HOST = "192.168.4.19"
IWORK_DB_PORT = 3306
IWORK_DB_NAME = "payroll"
CONNECT_TIMEOUT = 10
READ_TIMEOUT = 300

# ======
# 导出配置
CHUNK_SIZE = 10_000
SOURCE_FIELDS = (
    "TicketNo",
    "SeqNo",
    "WrkOrder",
    "BundleNo",
    "StepNo",
    "Qty",
    "RegPerSysID",
    "RegDate",
    "RegTime",
    "RFID",
    "Flow",
    "PO",
    "TimeCost",
    "SysSource",
    "AccBundleNo",
    "MtrType",
    "Color",
    "Sizx",
    "SerialNum",
    "StationID",
)
COMPACT_EXPORT_FIELDS = (
    "WrkOrder",
    "起始RegDate",
    *(f"StepNo={step_no} Qty总和" for step_no in COMPACT_STEP_NOS),
)
LOG_FORMAT = (
    "<green>{time:YYYY-MM-DD HH:mm:ss}</green> | "
    "<level>{level: <8}</level> | "
    "<cyan>{name}</cyan>:<cyan>{function}</cyan>:<cyan>{line}</cyan> - "
    "<level>{message}</level>"
)


def configure_logging() -> None:
    """配置控制台日志和保留七天的文件日志。"""
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    logger.remove()
    logger.add(sys.stdout, format=LOG_FORMAT, level="INFO")
    logger.add(
        LOG_DIR / "export_wrkorder_{time:YYYY-MM-DD}.log",
        format=LOG_FORMAT,
        rotation="00:00",
        retention="7 days",
        encoding="utf-8",
        level="DEBUG",
    )


def build_detail_query() -> str:
    """
    构建按完整工单号读取全部源明细的参数化查询。

    Returns:
        str: 只包含工单号过滤条件的只读查询。
    """
    fields = ", ".join(f"`{field}`" for field in SOURCE_FIELDS)
    return f"""
        SELECT {fields}
        FROM `pytckreg3`
        WHERE `WrkOrder` = %s
        ORDER BY `StepNo`, `RegDate`, `RegTime`, `TicketNo`, `SeqNo`
    """


def build_compact_query(wrkorder_count: int) -> str:
    """
    构建多工单精简汇总的参数化只读查询。

    最早 RegDate 从工单全部记录中获取，三个工序只限制各自的 Qty 汇总，
    因此不能在 WHERE 中提前过滤 StepNo。

    Args:
        wrkorder_count (int): 查询的工单数量。

    Returns:
        str: 按工单汇总最早日期和指定工序产量的查询。

    Raises:
        ValueError: 工单数量小于 1。
    """
    if wrkorder_count < 1:
        raise ValueError("精简模式至少需要一个 WRKORDER")

    placeholders = ", ".join("%s" for _ in range(wrkorder_count))
    qty_expressions = ",\n            ".join(
        "COALESCE(SUM(CASE WHEN `StepNo` = "
        f"{step_no} THEN `Qty` ELSE 0 END), 0) AS `StepNo_{step_no}_Qty`"
        for step_no in COMPACT_STEP_NOS
    )
    return f"""
        SELECT
            `WrkOrder`,
            MIN(`RegDate`) AS `StartRegDate`,
            {qty_expressions}
        FROM `pytckreg3`
        WHERE `WrkOrder` IN ({placeholders})
        GROUP BY `WrkOrder`
    """


def load_db_config() -> dict[str, Any]:
    """
    从现有 export.env 或 iwork/.env 加载生产库只读连接配置。

    Returns:
        dict[str, Any]: 可传给 pymysql.connect 的连接参数。

    Raises:
        ValueError: 未找到数据库用户名或密码。
    """
    values: dict[str, str] = {}
    for env_path in ENV_FILE_CANDIDATES:
        if not env_path.exists():
            continue
        for raw_line in env_path.read_text(encoding="utf-8").splitlines():
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            values[key.strip()] = value.strip()
        break

    user = values.get("IWORK_DB_USER", "")
    password = values.get("IWORK_DB_PASSWORD", "")
    if not user or not password:
        raise ValueError(
            "未找到生产库凭据，请在 scripts/export.env 中配置 "
            "IWORK_DB_USER 和 IWORK_DB_PASSWORD"
        )

    return {
        "host": values.get("IWORK_DB_HOST", IWORK_DB_HOST),
        "port": int(values.get("IWORK_DB_PORT", IWORK_DB_PORT)),
        "user": user,
        "password": password,
        "database": values.get("IWORK_DB_NAME", IWORK_DB_NAME),
        "charset": "utf8mb4",
        "connect_timeout": CONNECT_TIMEOUT,
        "read_timeout": READ_TIMEOUT,
        "cursorclass": SSCursor,
    }


def build_output_path(wrkorder: str, output_dir: Path) -> Path:
    """
    生成包含工单号和时间戳的导出文件路径。

    Args:
        wrkorder (str): 完整工单号。
        output_dir (Path): 导出目录。

    Returns:
        Path: 最终 XLSX 文件路径。
    """
    safe_wrkorder = re.sub(r'[^0-9A-Za-z_-]+', "_", wrkorder)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return output_dir / f"pytckreg3_{safe_wrkorder}_累计产量源明细_{timestamp}.xlsx"


def build_compact_output_path(output_dir: Path) -> Path:
    """
    生成精简汇总导出文件路径。

    Args:
        output_dir (Path): 导出目录。

    Returns:
        Path: 最终 XLSX 文件路径。
    """
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return output_dir / f"pytckreg3_工单累计产量精简汇总_{timestamp}.xlsx"


def export_wrkorder_details(
    wrkorder: str,
    output_dir: Path = OUTPUT_DIR,
    db_config: dict[str, Any] | None = None,
    connection_factory: Callable[..., Any] = pymysql.connect,
) -> tuple[Path, int]:
    """
    按完整工单号逐行导出生产表中的全部源数据。

    查询不限制日期、Flow 或工序号，也不做汇总、去重和字段转换。输出行按
    工序号和登记时间排序，便于直接核对每个工序的历史产量明细。

    Args:
        wrkorder (str): 要查询的完整 WrkOrder。
        output_dir (Path): XLSX 输出目录。
        db_config (dict[str, Any] | None): 数据库连接参数，默认从配置文件加载。
        connection_factory (Callable[..., Any]): 数据库连接工厂，默认使用 PyMySQL。

    Returns:
        tuple[Path, int]: 导出文件路径和源数据行数。

    Raises:
        ValueError: 工单号为空。
        pymysql.MySQLError: 生产库连接或查询失败。
    """
    normalized_wrkorder = wrkorder.strip()
    if not normalized_wrkorder:
        raise ValueError("WRKORDER 不能为空")

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = build_output_path(normalized_wrkorder, output_dir)
    temp_path = output_path.with_name(
        f"{output_path.stem}.{uuid4().hex}.tmp{output_path.suffix}"
    )
    connection = None
    workbook = None

    try:
        config = db_config if db_config is not None else load_db_config()
        logger.info(
            "开始查询生产源表: WrkOrder={}, host={}, database={}",
            normalized_wrkorder,
            config.get("host", ""),
            config.get("database", ""),
        )
        connection = connection_factory(**config)
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet(title="累计产量源明细")
        worksheet.append(SOURCE_FIELDS)

        row_count = 0
        with connection.cursor() as cursor:
            cursor.execute(build_detail_query(), (normalized_wrkorder,))
            for row in cursor:
                worksheet.append(row)
                row_count += 1
                if row_count % CHUNK_SIZE == 0:
                    logger.info("已写入 {} 条源数据", row_count)

        workbook.save(temp_path)
        workbook.close()
        workbook = None
        temp_path.replace(output_path)
        logger.success("导出完成: {}，共 {} 条源数据", output_path, row_count)
        return output_path, row_count
    except Exception:
        if workbook is not None:
            workbook.close()
        if temp_path.exists():
            temp_path.unlink()
        raise
    finally:
        if connection is not None:
            connection.close()


def export_compact_summary(
    wrkorders: tuple[str, ...],
    output_dir: Path = OUTPUT_DIR,
    db_config: dict[str, Any] | None = None,
    connection_factory: Callable[..., Any] = pymysql.connect,
) -> tuple[Path, int]:
    """
    批量导出各工单的最早日期和指定工序累计产量。

    查询结果会按传入工单顺序输出。生产表中不存在的工单仍保留一行，
    起始日期为空，三个工序的 Qty 总和为 0。

    Args:
        wrkorders (tuple[str, ...]): 要汇总的完整 WrkOrder 列表。
        output_dir (Path): XLSX 输出目录。
        db_config (dict[str, Any] | None): 数据库连接参数，默认从配置文件加载。
        connection_factory (Callable[..., Any]): 数据库连接工厂，默认使用 PyMySQL。

    Returns:
        tuple[Path, int]: 导出文件路径和输出工单行数。

    Raises:
        ValueError: 工单列表为空、包含空值或包含重复项。
        pymysql.MySQLError: 生产库连接或查询失败。
    """
    normalized_wrkorders = tuple(wrkorder.strip() for wrkorder in wrkorders)
    if not normalized_wrkorders:
        raise ValueError("精简模式至少需要一个 WRKORDER")
    if any(not wrkorder for wrkorder in normalized_wrkorders):
        raise ValueError("WRKORDER 列表不能包含空值")
    if len(set(normalized_wrkorders)) != len(normalized_wrkorders):
        raise ValueError("WRKORDER 列表不能包含重复项")

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = build_compact_output_path(output_dir)
    temp_path = output_path.with_name(
        f"{output_path.stem}.{uuid4().hex}.tmp{output_path.suffix}"
    )
    connection = None
    workbook = None

    try:
        config = db_config if db_config is not None else load_db_config()
        logger.info(
            "开始查询生产源表精简汇总: WRKORDER数量={}, host={}, database={}",
            len(normalized_wrkorders),
            config.get("host", ""),
            config.get("database", ""),
        )
        connection = connection_factory(**config)
        with connection.cursor() as cursor:
            cursor.execute(
                build_compact_query(len(normalized_wrkorders)),
                normalized_wrkorders,
            )
            query_rows = list(cursor)

        summaries = {str(row[0]): row[1:] for row in query_rows}
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet(title="工单累计产量精简汇总")
        worksheet.append(COMPACT_EXPORT_FIELDS)
        empty_summary = (None, *(0 for _ in COMPACT_STEP_NOS))
        for wrkorder in normalized_wrkorders:
            start_regdate, *qty_totals = summaries.get(wrkorder, empty_summary)
            worksheet.append(
                (
                    wrkorder,
                    start_regdate,
                    *(qty_total or 0 for qty_total in qty_totals),
                )
            )

        workbook.save(temp_path)
        workbook.close()
        workbook = None
        temp_path.replace(output_path)
        logger.success(
            "精简汇总导出完成: {}，输出 {} 个工单，其中 {} 个匹配生产数据",
            output_path,
            len(normalized_wrkorders),
            len(summaries),
        )
        return output_path, len(normalized_wrkorders)
    except Exception:
        if workbook is not None:
            workbook.close()
        if temp_path.exists():
            temp_path.unlink()
        raise
    finally:
        if connection is not None:
            connection.close()


def main() -> int:
    """
    根据全局模式执行完整明细或精简汇总导出。

    Returns:
        int: 成功返回 0，失败返回 1。
    """
    configure_logging()
    try:
        normalized_mode = EXPORT_MODE.strip().lower()
        if normalized_mode == "detail":
            export_wrkorder_details(WRKORDER)
        elif normalized_mode == "compact":
            export_compact_summary(WRKORDERS)
        else:
            raise ValueError(
                f"不支持的 EXPORT_MODE: {EXPORT_MODE}，只能使用 detail 或 compact"
            )
    except Exception as exc:
        logger.exception("导出失败: {}", exc)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
