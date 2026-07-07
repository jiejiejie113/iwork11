"""
导出 pytckreg3 表数据到 XLSX

表信息：
- 总行数: ~3900万条
- RegDate 有索引: PYTCKREG3_RegDate (可高效查询)

字段说明：
- RegDate: 登记日期（格式: 21/4/2026 00:00:00）
- RegTime: 登记时间（格式: 30/12/1899 16:25:53）

运行模式：
- export_date: 按 TARGET_DATE 导出单日全部数据
  输出: scripts/output/pytckreg3_20260421_143052.xlsx
- export_month: 按 TARGET_MONTH 导出数据（月份→整月合并单文件；日期→单日），
  仅保留 StepNo 在 STEPNO_FILTER 中的记录
  输出: scripts/output/pytckreg3_202602_step70_20260508_140539.xlsx
"""

import calendar
import sqlite3
import sys
import threading
from datetime import date, datetime
from pathlib import Path
from typing import Callable

import pymysql
from pymysql.cursors import SSCursor
from openpyxl import Workbook
from loguru import logger

logger.remove()
logger.add(
    sink=lambda msg: print(msg, end=""),
    format="<green>{time:HH:mm:ss}</green> | <level>{level: <8}</level> | <level>{message}</level>",
    level="INFO",
)

# ============== 可编辑参数 ==============
OUTPUT_DIR = Path(__file__).parent / "output"

# ---- 运行模式 ----
MODE = "export_date"            # "export_date" = 按日期导出 | "export_month" = 按月份导出

# ---- export_date 模式参数 ----
TARGET_DATE = "2026-06-25"      # None = 今日，或 格式"2026-04-21"

# ---- export_month 模式参数 ----
TARGET_MONTH = "2026-04"        # 支持月份 "2026-02" 或具体日期 "2026-02-20"，自动识别
STEPNO_FILTER = [70]            # StepNo 筛选值列表，如 [69] 或 [69, 70]

# ---- 数据库连接参数 ----
CONNECT_TIMEOUT = 10        # 连接超时（秒）
READ_TIMEOUT = 300          # 读取超时（秒）
CHUNK_SIZE = 10000          # 进度显示间隔
# IWORK_DB_HOST = "192.168.3.15"  # VCO
IWORK_DB_HOST = "192.168.4.19"  # EST

# ---- 数据库主机选项（供 GUI 使用） ----
DB_HOST_OPTIONS = {
    "VCO": "192.168.3.15",
    "EST": "192.168.4.19",
}
# ========================================


def get_db_config(host: str | None = None) -> dict:
    """
    获取数据库连接配置

    Args:
        host: 数据库主机地址，默认使用 IWORK_DB_HOST
    """
    # 按优先级查找专用 .env
    if getattr(sys, 'frozen', False):
        exe_dir = Path(sys.executable).parent
        env_paths = [
            exe_dir / "export.env",
            Path.cwd() / "export.env",
        ]
        # 调试：写入日志文件确认路径
        debug_log = exe_dir / "debug_paths.log"
        with open(debug_log, "w", encoding="utf-8") as dl:
            dl.write(f"sys.executable = {sys.executable}\n")
            dl.write(f"Path.cwd() = {Path.cwd()}\n")
            for i, p in enumerate(env_paths):
                dl.write(f"env_paths[{i}] = {p} (exists={p.exists()})\n")
    else:
        env_paths = [
            Path(__file__).parent / "export.env",
            Path.cwd() / "export.env",
            Path(__file__).parent.parent / "iwork" / ".env",
        ]

    config = {}
    for env_path in env_paths:
        if env_path.exists():
            with open(env_path, encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if line and "=" in line and not line.startswith("#"):
                        key, value = line.split("=", 1)
                        config[key.strip()] = value.strip()
            break

    return {
        "host": host or IWORK_DB_HOST,
        "port": int(config.get("IWORK_DB_PORT", 3306)),
        "user": config.get("IWORK_DB_USER", ""),
        "password": config.get("IWORK_DB_PASSWORD", ""),
        "database": config.get("IWORK_DB_NAME", ""),
        "charset": "utf8mb4",
        "connect_timeout": CONNECT_TIMEOUT,
        "read_timeout": READ_TIMEOUT,
    }


def _build_production_order_lookup() -> dict[str, tuple[str, str]]:
    """
    从 SQLite 读取生产工单数据，构建 style_no → (产品名称, 生产单号) 映射

    Returns:
        dict: {style_no: (product_name, order_no)}，文件不存在时返回空字典
    """
    sqlite_path = Path(__file__).parent.parent / "sqlite" / "production_orders.db"
    if not sqlite_path.exists():
        logger.warning(f"SQLite 文件不存在: {sqlite_path}，跳过工单信息匹配")
        return {}

    try:
        conn = sqlite3.connect(str(sqlite_path))
        cursor = conn.cursor()
        cursor.execute('SELECT "Style No", "Product Name", "order" FROM orders')
        lookup: dict[str, tuple[str, str]] = {}
        for row in cursor:
            style_no, product_name, order_no = row
            if style_no and style_no not in lookup:
                lookup[style_no] = (product_name or '', order_no or '')
        conn.close()
        logger.info(f"已加载 {len(lookup)} 条生产工单映射")
        return lookup
    except sqlite3.Error as e:
        logger.warning(f"SQLite 读取失败: {e}，跳过工单信息匹配")
        return {}


def export_to_xlsx(
    target_date: date,
    output_dir: Path,
    stepno_filter: list[int] | None = None,
    file_prefix: str = "pytckreg3",
    db_host: str | None = None,
    progress_callback: Callable[[int], None] | None = None,
    cancel_event: threading.Event | None = None,
) -> int:
    """
    导出数据到 XLSX

    Args:
        target_date: 目标日期
        output_dir: 输出目录
        stepno_filter: StepNo 筛选值列表，None 表示不过滤
        file_prefix: 输出文件名前缀
        db_host: 数据库主机地址，默认使用 IWORK_DB_HOST
        progress_callback: 进度回调，每 CHUNK_SIZE 行调用一次，参数为当前行数
        cancel_event: 取消事件，设置后中断导出

    Returns:
        int: 导出的记录数（取消时返回已导出行数的负值）
    """
    config = get_db_config(db_host)
    filter_info = f" | StepNo in {stepno_filter}" if stepno_filter else ""
    logger.info(f"连接: {config['host']}/{config['database']} | 日期: {target_date}{filter_info}")

    # 调试：记录连接信息（密码脱敏）
    if getattr(sys, 'frozen', False):
        exe_dir = Path(sys.executable).parent
        with open(exe_dir / "debug_paths.log", "a", encoding="utf-8") as dl:
            dl.write(f"\n--- export_to_xlsx ---\n")
            dl.write(f"host={config['host']}, port={config['port']}, db={config['database']}\n")
            dl.write(f"user={config['user']}, password={'***' if config['password'] else '(empty)'}\n")
            dl.write(f"target_date={target_date}, stepno_filter={stepno_filter}\n")

    try:
        conn = pymysql.connect(**config, cursorclass=SSCursor)
    except pymysql.Error as e:
        logger.error(f"连接失败: {e}")
        if getattr(sys, 'frozen', False):
            exe_dir = Path(sys.executable).parent
            with open(exe_dir / "debug_paths.log", "a", encoding="utf-8") as dl:
                dl.write(f"CONNECTION FAILED: {e}\n")
        return 0

    try:
        with conn.cursor() as cursor:
            sql = """
                SELECT TicketNo, SeqNo, WrkOrder, BundleNo, StepNo, Qty,
                       RegPerSysID, RegDate, RegTime, RFID, Flow, PO,
                       TimeCost, SysSource, AccBundleNo, MtrType, Color, Sizx,
                       SerialNum, StationID
                FROM pytckreg3
                WHERE RegDate >= %s AND RegDate < %s + INTERVAL 1 DAY
            """
            params: list = [target_date, target_date]

            if stepno_filter:
                placeholders = ",".join(["%s"] * len(stepno_filter))
                sql += f" AND StepNo IN ({placeholders})"
                params.extend(stepno_filter)

            cursor.execute(sql, params)

            columns = [d[0] for d in cursor.description]
            # 找到 WrkOrder 列的索引（用于后续匹配）
            wrk_order_idx = columns.index('WrkOrder')
            columns.append('ProductName')
            columns.append('OrderNo')

            lookup = _build_production_order_lookup()

            output_dir.mkdir(parents=True, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_file = output_dir / f"{file_prefix}_{timestamp}.xlsx"

            wb = Workbook(write_only=True)
            ws = wb.create_sheet()
            ws.append(columns)

            count = 0
            for row in cursor:
                if cancel_event and cancel_event.is_set():
                    logger.warning("用户取消导出")
                    wb.save(output_file)
                    wb.close()
                    logger.warning(f"已取消，部分数据已保存: {output_file} ({count} 条)")
                    return -count

                wrk_order = row[wrk_order_idx] or ''
                prefix = wrk_order[:6] if len(wrk_order) >= 6 else ''
                product_name, order_no = lookup.get(prefix, ('', ''))
                ws.append(row + (product_name, order_no))
                count += 1
                if count % CHUNK_SIZE == 0:
                    logger.info(f"已写入 {count} 条")
                    if progress_callback:
                        progress_callback(count)

            wb.save(output_file)
            wb.close()
            logger.success(f"导出完成: {output_file} ({count} 条)")
            if getattr(sys, 'frozen', False):
                exe_dir = Path(sys.executable).parent
                with open(exe_dir / "debug_paths.log", "a", encoding="utf-8") as dl:
                    dl.write(f"SUCCESS: {count} 条, file={output_file}\n")
            return count

    except pymysql.Error as e:
        logger.error(f"查询失败: {e}")
        if getattr(sys, 'frozen', False):
            exe_dir = Path(sys.executable).parent
            with open(exe_dir / "debug_paths.log", "a", encoding="utf-8") as dl:
                dl.write(f"QUERY FAILED: {e}\n")
        return 0
    finally:
        try:
            conn.close()
        except Exception:
            pass


def export_month_to_xlsx(
    target_month: str,
    output_dir: Path,
    stepno_filter: list[int] | None = None,
    db_host: str | None = None,
    progress_callback: Callable[[int], None] | None = None,
    cancel_event: threading.Event | None = None,
) -> int:
    """
    导出数据（自动识别月份或日期），仅保留 StepNo 在筛选列表中的记录

    - "YYYY-MM"（如 "2026-02"）：导出整月，合并为单个 XLSX 文件
    - "YYYY-MM-DD"（如 "2026-02-20"）：导出单日

    Args:
        target_month: 目标月份 "YYYY-MM" 或具体日期 "YYYY-MM-DD"
        stepno_filter: StepNo 筛选值列表
        output_dir: 输出目录
        db_host: 数据库主机地址
        progress_callback: 进度回调
        cancel_event: 取消事件

    Returns:
        int: 导出的总记录数（取消时为负数）
    """
    parts = target_month.split("-")
    if stepno_filter:
        stepno_str = "_".join(str(s) for s in stepno_filter)
        file_suffix = f"_step{stepno_str}"
    else:
        stepno_str = "all"
        file_suffix = ""

    if len(parts) == 3:
        target_date = datetime.strptime(target_month, "%Y-%m-%d").date()
        logger.info(f"日期导出模式: {target_date}, StepNo in {stepno_filter}")

        return export_to_xlsx(
            target_date=target_date,
            output_dir=output_dir,
            stepno_filter=stepno_filter,
            file_prefix=f"pytckreg3_{target_date.strftime('%Y%m%d')}{file_suffix}",
            db_host=db_host,
            progress_callback=progress_callback,
            cancel_event=cancel_event,
        )

    elif len(parts) == 2:
        year, month = map(int, parts)
        _, last_day = calendar.monthrange(year, month)
        start_date = date(year, month, 1)
        end_date = date(year, month, last_day)

        logger.info(f"月份导出模式: {target_month} ({start_date} ~ {end_date}), StepNo in {stepno_filter}")

        config = get_db_config(db_host)
        try:
            conn = pymysql.connect(**config, cursorclass=SSCursor)
        except pymysql.Error as e:
            logger.error(f"连接失败: {e}")
            raise

        try:
            with conn.cursor() as cursor:
                cursor.execute(
                    "SELECT TicketNo, SeqNo, WrkOrder, BundleNo, StepNo, Qty, "
                    "RegPerSysID, RegDate, RegTime, RFID, Flow, PO, "
                    "TimeCost, SysSource, AccBundleNo, MtrType, Color, Sizx, "
                    "SerialNum, StationID FROM pytckreg3 LIMIT 0"
                )
                columns = [d[0] for d in cursor.description]
                wrk_order_idx = columns.index('WrkOrder')
                columns.append('ProductName')
                columns.append('OrderNo')

                lookup = _build_production_order_lookup()

                output_dir.mkdir(parents=True, exist_ok=True)
                month_str = target_month.replace("-", "")
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                output_file = output_dir / f"pytckreg3_{month_str}{file_suffix}_{timestamp}.xlsx"

                wb = Workbook(write_only=True)
                ws = wb.create_sheet()
                ws.append(columns)

                total = 0
                current = start_date
                while current <= end_date:
                    if cancel_event and cancel_event.is_set():
                        logger.warning("用户取消导出")
                        wb.save(output_file)
                        wb.close()
                        logger.warning(f"已取消，部分数据已保存: {output_file} ({total} 条)")
                        return -total

                    sql = """
                        SELECT TicketNo, SeqNo, WrkOrder, BundleNo, StepNo, Qty,
                               RegPerSysID, RegDate, RegTime, RFID, Flow, PO,
                               TimeCost, SysSource, AccBundleNo, MtrType, Color, Sizx,
                               SerialNum, StationID
                        FROM pytckreg3
                        WHERE RegDate >= %s AND RegDate < %s + INTERVAL 1 DAY
                    """
                    params: list = [current, current]

                    if stepno_filter:
                        placeholders = ",".join(["%s"] * len(stepno_filter))
                        sql += f" AND StepNo IN ({placeholders})"
                        params.extend(stepno_filter)

                    cursor.execute(sql, params)
                    daily_count = 0
                    for row in cursor:
                        if cancel_event and cancel_event.is_set():
                            wb.save(output_file)
                            wb.close()
                            logger.warning(f"已取消，部分数据已保存: {output_file} ({total} 条)")
                            return -total

                        wrk_order = row[wrk_order_idx] or ''
                        prefix = wrk_order[:6] if len(wrk_order) >= 6 else ''
                        product_name, order_no = lookup.get(prefix, ('', ''))
                        ws.append(row + (product_name, order_no))
                        total += 1
                        daily_count += 1
                        if total % CHUNK_SIZE == 0:
                            logger.info(f"已写入 {total} 条")
                            if progress_callback:
                                progress_callback(total)

                    logger.info(f"  {current}: {daily_count} 条 (累计 {total})")
                    current = date.fromordinal(current.toordinal() + 1)

                wb.save(output_file)
                wb.close()
                logger.success(f"月份导出完成: {output_file} ({total} 条)")
                return total

        except pymysql.Error as e:
            logger.error(f"查询失败: {e}")
            return 0
        finally:
            try:
                conn.close()
            except Exception:
                pass

    else:
        logger.error(f"无法识别的格式: {target_month}，请使用 YYYY-MM 或 YYYY-MM-DD")
        return 0


def main():
    if MODE == "export_date":
        if TARGET_DATE:
            target_date = datetime.strptime(TARGET_DATE, "%Y-%m-%d").date()
        else:
            target_date = date.today()
        export_to_xlsx(target_date, OUTPUT_DIR)

    elif MODE == "export_month":
        export_month_to_xlsx(TARGET_MONTH, OUTPUT_DIR, STEPNO_FILTER)

    else:
        logger.error(f"未知模式: {MODE}，可选值: export_date / export_month")


if __name__ == "__main__":
    main()
